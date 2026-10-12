from pathlib import Path
from typing import List, Optional, Any
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models.irsaliye import DocumentItem, IrsaliyeData
from ..models.analytics import AnalyticsResponse
from ..core.config import settings

class ExcelExporter:
    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SUBHEADER_FILL = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    CARD_HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    ACCENT_FILL = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    LIGHT_BLUE_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    BORDER_COLOR = Side(border_style="thin", color="CBD5E1")
    THIN_BORDER = Border(left=BORDER_COLOR, right=BORDER_COLOR, top=BORDER_COLOR, bottom=BORDER_COLOR)

    @staticmethod
    def sanitize_cell_value(val: Any) -> Any:
        """
        Formül Enjeksiyonu (CSV/Excel Formula Injection) koruması.
        Hücre içeriği '=', '+', '-', '@', tab veya cr ile başlıyorsa başına tek tırnak (') ekler.
        """
        if isinstance(val, str) and len(val) > 0:
            if val[0] in ('=', '+', '-', '@', '\t', '\r'):
                return "'" + val
        return val

    @classmethod
    def export_single_document(cls, doc: DocumentItem, output_path: Optional[Path] = None) -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "İrsaliye Detayı"
        ws.views.sheetView[0].showGridLines = True

        data = doc.veri
        if not data:
            raise ValueError("Belgede ayrıştırılmış veri bulunmuyor.")

        # Başlık Blok
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = "SEVK İRSALİYESİ AYRIŞTIRMA VE DOĞRULAMA RAPORU"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Üst Bilgiler Tablosu
        headers_info = [
            ("İrsaliye No:", cls.sanitize_cell_value(data.header.irsaliye_no or "-"), "Düzenleme Tarihi:", cls.sanitize_cell_value(data.header.duzenleme_tarihi or "-")),
            ("Belge Türü:", cls.sanitize_cell_value(data.header.irsaliye_tipi or "-"), "Fiili Sevk Tarihi:", cls.sanitize_cell_value(data.header.sevk_tarihi or "-")),
            ("Gönderici Ünvanı:", cls.sanitize_cell_value(data.header.gonderici_unvan or "-"), "Sevk Saati:", cls.sanitize_cell_value(data.header.sevk_saati or "-")),
            ("Gönderici VKN:", cls.sanitize_cell_value(data.header.gonderici_vkn or "-"), "Alıcı Ünvanı:", cls.sanitize_cell_value(data.header.alici_unvan or "-")),
            ("Gönderici Adresi:", cls.sanitize_cell_value(data.header.gonderici_adres or "-"), "Alıcı VKN:", cls.sanitize_cell_value(data.header.alici_vkn or "-")),
            ("Taşıyıcı / Ambar:", cls.sanitize_cell_value(data.header.tasiyici_unvan or "-"), "Teslim Adresi:", cls.sanitize_cell_value(data.header.teslim_adresi or "-")),
            ("Araç Plakası:", cls.sanitize_cell_value(data.header.plaka or "-"), "Şoför:", cls.sanitize_cell_value(data.header.sofor_ad_soyad or "-")),
        ]

        row_idx = 3
        for label1, val1, label2, val2 in headers_info:
            ws.cell(row=row_idx, column=1, value=label1).font = Font(name="Segoe UI", size=10, bold=True, color="475569")
            ws.cell(row=row_idx, column=2, value=val1).font = Font(name="Segoe UI", size=10, bold=False)
            ws.cell(row=row_idx, column=4, value=label2).font = Font(name="Segoe UI", size=10, bold=True, color="475569")
            ws.cell(row=row_idx, column=5, value=val2).font = Font(name="Segoe UI", size=10, bold=False)
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1

        # Kalemler Başlığı
        row_idx += 1
        ws.merge_cells(f"A{row_idx}:F{row_idx}")
        items_header = ws[f"A{row_idx}"]
        items_header.value = "SEVK EDİLEN MAL / HİZMET KALEMLERİ"
        items_header.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        items_header.fill = cls.SUBHEADER_FILL
        items_header.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row_idx].height = 25

        # Kalem Tablo Sütunları
        row_idx += 1
        table_cols = ["Sıra No", "Mal / Malzeme Cinsi ve Açıklaması", "Miktar", "Birim", "Birim Fiyat", "Tutar / Notlar"]
        for c_idx, col_name in enumerate(table_cols, start=1):
            cell = ws.cell(row=row_idx, column=c_idx, value=col_name)
            cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 4) else "left", vertical="center")
            cell.border = cls.THIN_BORDER
        ws.row_dimensions[row_idx].height = 24

        # Satırlar
        total_qty = 0.0
        for item in data.kalemler:
            row_idx += 1
            c1 = ws.cell(row=row_idx, column=1, value=item.sira_no or "")
            c2 = ws.cell(row=row_idx, column=2, value=cls.sanitize_cell_value(item.mal_hizmet or ""))
            c3 = ws.cell(row=row_idx, column=3, value=item.miktar if item.miktar is not None else "")
            c4 = ws.cell(row=row_idx, column=4, value=cls.sanitize_cell_value(item.birim or ""))
            c5 = ws.cell(row=row_idx, column=5, value=item.birim_fiyat if item.birim_fiyat is not None else "")
            c6 = ws.cell(row=row_idx, column=6, value=cls.sanitize_cell_value(item.notlar or ""))

            c1.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="right")
            c4.alignment = Alignment(horizontal="center")
            c5.alignment = Alignment(horizontal="right")

            for cell in (c1, c2, c3, c4, c5, c6):
                cell.font = Font(name="Segoe UI", size=10)
                cell.border = cls.THIN_BORDER

            if item.miktar:
                total_qty += float(item.miktar)
            ws.row_dimensions[row_idx].height = 20

        # Toplam Satırı
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="TOPLAM MİKTAR:").font = Font(name="Segoe UI", size=10, bold=True)
        tot_cell = ws.cell(row=row_idx, column=3, value=total_qty)
        tot_cell.font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
        tot_cell.alignment = Alignment(horizontal="right")

        # Notlar
        if data.genel_notlar:
            row_idx += 2
            ws.cell(row=row_idx, column=1, value="El Yazısı / Ambar Teslim Notları:").font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
            ws.cell(row=row_idx, column=2, value=cls.sanitize_cell_value(data.genel_notlar)).font = Font(name="Segoe UI", size=10, italic=True)

        cls._adjust_column_widths(ws)

        target_file = output_path or (settings.EXPORT_DIR / f"irsaliye_{data.header.irsaliye_no or doc.id}.xlsx")
        wb.save(target_file)
        return target_file

    @classmethod
    def export_batch_summary(cls, documents: List[DocumentItem], output_path: Optional[Path] = None) -> Path:
        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = "Tüm İrsaliyeler Özeti"
        ws_summary.views.sheetView[0].showGridLines = True

        ws_summary.merge_cells("A1:K1")
        title_cell = ws_summary["A1"]
        title_cell.value = "TOPLU İRSALİYE VE SEVKİYAT İŞLEME RAPORU"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_summary.row_dimensions[1].height = 32

        summary_headers = [
            "Dosya Adı", "İrsaliye No", "Düzenleme Tarihi", "Sevk Tarihi",
            "Gönderici Ünvan", "Gönderici VKN", "Alıcı Ünvan", "Alıcı VKN",
            "Araç Plakası", "Kalem Sayısı", "Durum"
        ]
        for c_idx, h_text in enumerate(summary_headers, start=1):
            c = ws_summary.cell(row=2, column=c_idx, value=h_text)
            c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            c.fill = cls.SUBHEADER_FILL
            c.alignment = Alignment(horizontal="center" if c_idx in (2, 3, 4, 6, 8, 9, 10, 11) else "left", vertical="center")
            c.border = cls.THIN_BORDER
        ws_summary.row_dimensions[2].height = 24

        ws_items = wb.create_sheet(title="Tüm Ürün Kalemleri")
        ws_items.views.sheetView[0].showGridLines = True
        item_headers = [
            "İrsaliye No", "Tarih", "Gönderici", "Alıcı",
            "Sıra No", "Mal / Malzeme Cinsi", "Miktar", "Birim", "Notlar"
        ]
        for c_idx, h_text in enumerate(item_headers, start=1):
            c = ws_items.cell(row=1, column=c_idx, value=h_text)
            c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            c.fill = cls.HEADER_FILL
            c.alignment = Alignment(horizontal="center" if c_idx in (1, 2, 5, 7, 8) else "left", vertical="center")
            c.border = cls.THIN_BORDER
        ws_items.row_dimensions[1].height = 24

        s_row = 3
        i_row = 2

        for doc in documents:
            if not doc.veri:
                continue
            v = doc.veri
            h = v.header

            ws_summary.cell(row=s_row, column=1, value=cls.sanitize_cell_value(doc.dosya_adi))
            ws_summary.cell(row=s_row, column=2, value=cls.sanitize_cell_value(h.irsaliye_no or "-"))
            ws_summary.cell(row=s_row, column=3, value=cls.sanitize_cell_value(h.duzenleme_tarihi or "-"))
            ws_summary.cell(row=s_row, column=4, value=cls.sanitize_cell_value(h.sevk_tarihi or "-"))
            ws_summary.cell(row=s_row, column=5, value=cls.sanitize_cell_value(h.gonderici_unvan or "-"))
            ws_summary.cell(row=s_row, column=6, value=cls.sanitize_cell_value(h.gonderici_vkn or "-"))
            ws_summary.cell(row=s_row, column=7, value=cls.sanitize_cell_value(h.alici_unvan or "-"))
            ws_summary.cell(row=s_row, column=8, value=cls.sanitize_cell_value(h.alici_vkn or "-"))
            ws_summary.cell(row=s_row, column=9, value=cls.sanitize_cell_value(h.plaka or "-"))
            ws_summary.cell(row=s_row, column=10, value=len(v.kalemler))
            ws_summary.cell(row=s_row, column=11, value=doc.durum.value)

            for col in range(1, 12):
                cell = ws_summary.cell(row=s_row, column=col)
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = cls.THIN_BORDER
            ws_summary.row_dimensions[s_row].height = 20
            s_row += 1

            for item in v.kalemler:
                ws_items.cell(row=i_row, column=1, value=cls.sanitize_cell_value(h.irsaliye_no or "-"))
                ws_items.cell(row=i_row, column=2, value=cls.sanitize_cell_value(h.duzenleme_tarihi or "-"))
                ws_items.cell(row=i_row, column=3, value=cls.sanitize_cell_value(h.gonderici_unvan or "-"))
                ws_items.cell(row=i_row, column=4, value=cls.sanitize_cell_value(h.alici_unvan or "-"))
                ws_items.cell(row=i_row, column=5, value=item.sira_no or "")
                ws_items.cell(row=i_row, column=6, value=cls.sanitize_cell_value(item.mal_hizmet or ""))
                ws_items.cell(row=i_row, column=7, value=item.miktar if item.miktar is not None else "")
                ws_items.cell(row=i_row, column=8, value=cls.sanitize_cell_value(item.birim or ""))
                ws_items.cell(row=i_row, column=9, value=cls.sanitize_cell_value(item.notlar or ""))

                for col in range(1, 10):
                    cell = ws_items.cell(row=i_row, column=col)
                    cell.font = Font(name="Segoe UI", size=9)
                    cell.border = cls.THIN_BORDER
                ws_items.row_dimensions[i_row].height = 20
                i_row += 1

        cls._adjust_column_widths(ws_summary)
        cls._adjust_column_widths(ws_items)

        target_file = output_path or (settings.EXPORT_DIR / "tum_irsaliyeler_toplu_rapor.xlsx")
        wb.save(target_file)
        return target_file

    @classmethod
    def export_executive_report(
        cls,
        analytics_data: AnalyticsResponse,
        documents: List[DocumentItem],
        output_path: Optional[Path] = None
    ) -> Path:
        """
        Lojistik Yönetici Kokpiti Çok Sayfalı Kurumsal Excel Raporu üretir.
        Sayfalar:
          1. Yonetici_Ozeti (Üst düzey KPI tablosu ve operasyonel sağlık)
          2. Tedarikci_Dagilimi (Tüm tedarikçilerin hacim sıralaması ve payları)
          3. En_Cok_Gelen_Kalemler (Mal/malzeme kalemleri miktar ve frekans sıralaması)
          4. Evrak_Dokumu (Tüm irsaliyelerin ayrıntılı işlem ve uyuşmazlık dökümü)
        """
        wb = openpyxl.Workbook()

        # ==========================================
        # 1. SAYFA: Yonetici_Ozeti
        # ==========================================
        ws_exec = wb.active
        ws_exec.title = "Yonetici_Ozeti"
        ws_exec.views.sheetView[0].showGridLines = True

        # Ana Başlık Banner
        ws_exec.merge_cells("A1:G1")
        title_cell = ws_exec["A1"]
        title_cell.value = "DOCULOG WORKSTATION — RESMİ MAL KABUL VE LOJİSTİK YÖNETİCİ RAPORU"
        title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_exec.row_dimensions[1].height = 36

        # Meta Bilgi Satırları
        meta_info = [
            ("Rapor Tarihi:", analytics_data.son_guncelleme or datetime.now().strftime("%d.%m.%Y %H:%M:%S")),
            ("Raporu Üreten Sistem:", "DocuLog Workstation Enterprise Logistics Cockpit"),
            ("Genel Mal Kabul Durumu:", "Aktif Sevkiyat Akışı" if not analytics_data.empty else "Evrak Kuyruğu Boş")
        ]
        r_meta = 3
        for label, val in meta_info:
            ws_exec.cell(row=r_meta, column=1, value=label).font = Font(name="Segoe UI", size=9, bold=True, color="475569")
            ws_exec.cell(row=r_meta, column=2, value=cls.sanitize_cell_value(val)).font = Font(name="Segoe UI", size=9)
            ws_exec.row_dimensions[r_meta].height = 18
            r_meta += 1

        # KPI Bölüm Başlığı
        r_kpi = r_meta + 1
        ws_exec.merge_cells(f"A{r_kpi}:C{r_kpi}")
        kpi_hdr = ws_exec[f"A{r_kpi}"]
        kpi_hdr.value = "1. TEMEL OPERASYONEL KPI GÖSTERGELERİ"
        kpi_hdr.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        kpi_hdr.fill = cls.SUBHEADER_FILL
        kpi_hdr.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_exec.row_dimensions[r_kpi].height = 24

        kpi = analytics_data.kpi
        kpi_rows = [
            ("Toplam İrsaliye Hacmi", kpi.toplam_irsaliye_hacmi, "Evrak"),
            ("Mal Kabul Onay Oranı", f"%{kpi.onay_orani_yuzde}", "Onaylanan / (Onaylanan + Tamamlanan + Hata)"),
            ("Onaylanan Belge Adedi", kpi.onaylanan_belge, "Evrak (Doğrulandı & Kabul Edildi)"),
            ("Tamamlanan / İşlenen Belge", kpi.tamamlanan_belge, "Evrak (OCR Başarılı)"),
            ("Toplam İşlenen Kalem Adedi", kpi.toplam_islenen_kalem_adedi, "Satır Kalemi"),
            ("Toplam İşlenen Mal Miktarı", kpi.toplam_islenen_miktar, "Sayısal Birim Toplamı"),
            ("Ortalama Dijitalleştirme Hızı", f"{kpi.ortalama_dijitallesme_hizi_sn} sn", "Evrak Başına Yapay Zeka Hızı")
        ]

        r_kpi += 1
        ws_exec.cell(row=r_kpi, column=1, value="Metrik Göstergesi").font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws_exec.cell(row=r_kpi, column=1).fill = cls.CARD_HEADER_FILL
        ws_exec.cell(row=r_kpi, column=1).border = cls.THIN_BORDER

        ws_exec.cell(row=r_kpi, column=2, value="Değer").font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws_exec.cell(row=r_kpi, column=2).fill = cls.CARD_HEADER_FILL
        ws_exec.cell(row=r_kpi, column=2).alignment = Alignment(horizontal="right")
        ws_exec.cell(row=r_kpi, column=2).border = cls.THIN_BORDER

        ws_exec.cell(row=r_kpi, column=3, value="Birim / Metodoloji").font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws_exec.cell(row=r_kpi, column=3).fill = cls.CARD_HEADER_FILL
        ws_exec.cell(row=r_kpi, column=3).border = cls.THIN_BORDER
        ws_exec.row_dimensions[r_kpi].height = 22

        for label, val, note in kpi_rows:
            r_kpi += 1
            c1 = ws_exec.cell(row=r_kpi, column=1, value=label)
            c2 = ws_exec.cell(row=r_kpi, column=2, value=cls.sanitize_cell_value(val))
            c3 = ws_exec.cell(row=r_kpi, column=3, value=cls.sanitize_cell_value(note))

            c1.font = Font(name="Segoe UI", size=9, bold=True)
            c2.font = Font(name="Segoe UI", size=9, bold=True, color="1E3A8A")
            c2.alignment = Alignment(horizontal="right")
            c3.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

            for c in (c1, c2, c3):
                c.border = cls.THIN_BORDER
            ws_exec.row_dimensions[r_kpi].height = 20

        # Operasyonel Durum ve Uyuşmazlıklar Bölümü
        r_stat = r_kpi + 2
        ws_exec.merge_cells(f"A{r_stat}:C{r_stat}")
        stat_hdr = ws_exec[f"A{r_stat}"]
        stat_hdr.value = "2. OPERASYONEL DURUM VE UYUŞMAZLIK ANALİZİ"
        stat_hdr.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        stat_hdr.fill = cls.SUBHEADER_FILL
        stat_hdr.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_exec.row_dimensions[r_stat].height = 24

        durum = analytics_data.durum_dagilimi
        durum_rows = [
            ("Onaylandı (Kabul Edildi)", durum.onaylandi, "Sevkiyat eksiksiz onaylandı"),
            ("Tamamlandı (Onay Bekleyen)", durum.tamamlandi, "Ayrıştırma yapıldı, kabul onayı bekliyor"),
            ("Kuyrukta Bekleyen", durum.kuyrukta, "İşleme sırası bekleyen evraklar"),
            ("İşlenmekte Olan", durum.isleniyor, "Yapay zeka analiz sürecinde"),
            ("Hata Alan Belge", durum.hata, "Okuma veya format hatası"),
            ("Tespit Edilen Uyuşmazlık Adedi", durum.uyusmazlik_sayisi, "Hasar, eksik, miktar farkı veya düşük güven"),
            ("Operasyonel Uyuşmazlık Oranı", f"%{durum.uyusmazlik_orani_yuzde}", "Uyuşmazlıkların toplam hacme oranı"),
            ("Ortalama OCR Güven Skoru", f"{durum.ortalama_guven_skoru:.2f}" if durum.ortalama_guven_skoru else "0.00", "Model güven düzeyi (0.00 - 1.00)")
        ]

        r_stat += 1
        ws_exec.cell(row=r_stat, column=1, value="Operasyonel Durum").font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws_exec.cell(row=r_stat, column=1).fill = cls.CARD_HEADER_FILL
        ws_exec.cell(row=r_stat, column=1).border = cls.THIN_BORDER

        ws_exec.cell(row=r_stat, column=2, value="Adet / Değer").font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws_exec.cell(row=r_stat, column=2).fill = cls.CARD_HEADER_FILL
        ws_exec.cell(row=r_stat, column=2).alignment = Alignment(horizontal="right")
        ws_exec.cell(row=r_stat, column=2).border = cls.THIN_BORDER

        ws_exec.cell(row=r_stat, column=3, value="Açıklama").font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws_exec.cell(row=r_stat, column=3).fill = cls.CARD_HEADER_FILL
        ws_exec.cell(row=r_stat, column=3).border = cls.THIN_BORDER
        ws_exec.row_dimensions[r_stat].height = 22

        for label, val, note in durum_rows:
            r_stat += 1
            c1 = ws_exec.cell(row=r_stat, column=1, value=label)
            c2 = ws_exec.cell(row=r_stat, column=2, value=cls.sanitize_cell_value(val))
            c3 = ws_exec.cell(row=r_stat, column=3, value=cls.sanitize_cell_value(note))

            c1.font = Font(name="Segoe UI", size=9)
            c2.font = Font(name="Segoe UI", size=9, bold=True)
            c2.alignment = Alignment(horizontal="right")
            c3.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

            for c in (c1, c2, c3):
                c.border = cls.THIN_BORDER
            ws_exec.row_dimensions[r_stat].height = 20

        cls._adjust_column_widths(ws_exec)

        # ==========================================
        # 2. SAYFA: Tedarikci_Dagilimi
        # ==========================================
        ws_sup = wb.create_sheet(title="Tedarikci_Dagilimi")
        ws_sup.views.sheetView[0].showGridLines = True

        ws_sup.merge_cells("A1:G1")
        sup_title = ws_sup["A1"]
        sup_title.value = "TEDARİKÇİ VE CARİ SEVKİYAT HACİM DAĞILIMI"
        sup_title.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        sup_title.fill = cls.HEADER_FILL
        sup_title.alignment = Alignment(horizontal="center", vertical="center")
        ws_sup.row_dimensions[1].height = 32

        sup_headers = ["Sıra", "Tedarikçi Ünvanı", "VKN / TCKN", "Evrak Sayısı", "Toplam Kalem", "Toplam Mal Miktarı", "Hacim Payı (%)"]
        for c_idx, h_text in enumerate(sup_headers, start=1):
            cell = ws_sup.cell(row=2, column=c_idx, value=h_text)
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            cell.fill = cls.SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 4, 5, 7) else "left", vertical="center")
            cell.border = cls.THIN_BORDER
        ws_sup.row_dimensions[2].height = 24

        r_s = 3
        total_sup_docs = 0
        total_sup_items = 0
        total_sup_qty = 0.0

        for idx, sup in enumerate(analytics_data.tedarikci_dagilimi, start=1):
            c1 = ws_sup.cell(row=r_s, column=1, value=idx)
            c2 = ws_sup.cell(row=r_s, column=2, value=cls.sanitize_cell_value(sup.unvan))
            c3 = ws_sup.cell(row=r_s, column=3, value=cls.sanitize_cell_value(sup.vkn or "-"))
            c4 = ws_sup.cell(row=r_s, column=4, value=sup.evrak_sayisi)
            c5 = ws_sup.cell(row=r_s, column=5, value=sup.kalem_sayisi)
            c6 = ws_sup.cell(row=r_s, column=6, value=sup.toplam_miktar)
            c7 = ws_sup.cell(row=r_s, column=7, value=f"%{sup.pay_yuzde}")

            c1.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="center")
            c4.alignment = Alignment(horizontal="center")
            c5.alignment = Alignment(horizontal="center")
            c6.alignment = Alignment(horizontal="right")
            c7.alignment = Alignment(horizontal="right")

            use_fill = cls.ZEBRA_FILL if idx % 2 == 0 else None
            for cell in (c1, c2, c3, c4, c5, c6, c7):
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = cls.THIN_BORDER
                if use_fill:
                    cell.fill = use_fill

            total_sup_docs += sup.evrak_sayisi
            total_sup_items += sup.kalem_sayisi
            total_sup_qty += sup.toplam_miktar
            ws_sup.row_dimensions[r_s].height = 20
            r_s += 1

        # Genel Toplam Satırı
        if analytics_data.tedarikci_dagilimi:
            ws_sup.cell(row=r_s, column=2, value="GENEL TOPLAM:").font = Font(name="Segoe UI", size=9, bold=True)
            ws_sup.cell(row=r_s, column=4, value=total_sup_docs).font = Font(name="Segoe UI", size=9, bold=True)
            ws_sup.cell(row=r_s, column=4).alignment = Alignment(horizontal="center")
            ws_sup.cell(row=r_s, column=5, value=total_sup_items).font = Font(name="Segoe UI", size=9, bold=True)
            ws_sup.cell(row=r_s, column=5).alignment = Alignment(horizontal="center")
            ws_sup.cell(row=r_s, column=6, value=round(total_sup_qty, 2)).font = Font(name="Segoe UI", size=9, bold=True, color="1E3A8A")
            ws_sup.cell(row=r_s, column=6).alignment = Alignment(horizontal="right")
            ws_sup.cell(row=r_s, column=7, value="%100.0").font = Font(name="Segoe UI", size=9, bold=True)
            ws_sup.cell(row=r_s, column=7).alignment = Alignment(horizontal="right")

            for c_idx in range(1, 8):
                ws_sup.cell(row=r_s, column=c_idx).border = cls.THIN_BORDER
                ws_sup.cell(row=r_s, column=c_idx).fill = cls.LIGHT_BLUE_FILL
            ws_sup.row_dimensions[r_s].height = 22

        cls._adjust_column_widths(ws_sup)

        # ==========================================
        # 3. SAYFA: En_Cok_Gelen_Kalemler
        # ==========================================
        ws_top = wb.create_sheet(title="En_Cok_Gelen_Kalemler")
        ws_top.views.sheetView[0].showGridLines = True

        ws_top.merge_cells("A1:E1")
        top_title = ws_top["A1"]
        top_title.value = "EN ÇOK HAREKET GÖREN ÜRÜN VE MALZEME KALEMLERİ"
        top_title.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        top_title.fill = cls.HEADER_FILL
        top_title.alignment = Alignment(horizontal="center", vertical="center")
        ws_top.row_dimensions[1].height = 32

        top_headers = ["Sıra", "Mal / Malzeme Cinsi ve Açıklaması", "Geliş Sıklığı (Frekans)", "Ölçü Birimi", "Toplam Miktar"]
        for c_idx, h_text in enumerate(top_headers, start=1):
            cell = ws_top.cell(row=2, column=c_idx, value=h_text)
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            cell.fill = cls.SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 4) else "left", vertical="center")
            cell.border = cls.THIN_BORDER
        ws_top.row_dimensions[2].height = 24

        r_t = 3
        for idx, item in enumerate(analytics_data.en_cok_gelen_kalemler, start=1):
            c1 = ws_top.cell(row=r_t, column=1, value=idx)
            c2 = ws_top.cell(row=r_t, column=2, value=cls.sanitize_cell_value(item.mal_hizmet))
            c3 = ws_top.cell(row=r_t, column=3, value=item.frekans)
            c4 = ws_top.cell(row=r_t, column=4, value=cls.sanitize_cell_value(item.birim))
            c5 = ws_top.cell(row=r_t, column=5, value=item.toplam_miktar)

            c1.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="center")
            c4.alignment = Alignment(horizontal="center")
            c5.alignment = Alignment(horizontal="right")

            use_fill = cls.ZEBRA_FILL if idx % 2 == 0 else None
            for cell in (c1, c2, c3, c4, c5):
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = cls.THIN_BORDER
                if use_fill:
                    cell.fill = use_fill
            ws_top.row_dimensions[r_t].height = 20
            r_t += 1

        cls._adjust_column_widths(ws_top)

        # ==========================================
        # 4. SAYFA: Evrak_Dokumu
        # ==========================================
        ws_docs = wb.create_sheet(title="Evrak_Dokumu")
        ws_docs.views.sheetView[0].showGridLines = True

        ws_docs.merge_cells("A1:K1")
        docs_title = ws_docs["A1"]
        docs_title.value = "DETAYLI İRSALİYE VE SEVKİYAT EVRAK DÖKÜMÜ"
        docs_title.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        docs_title.fill = cls.HEADER_FILL
        docs_title.alignment = Alignment(horizontal="center", vertical="center")
        ws_docs.row_dimensions[1].height = 32

        doc_headers = [
            "İrsaliye No", "Düzenleme Tarihi", "Gönderici Ünvan", "Gönderici VKN",
            "Alıcı Ünvan", "Araç Plakası", "Kalem Sayısı", "Toplam Miktar",
            "OCR Güven Skoru", "İşlem Durumu", "Ambar Notu / Uyuşmazlık Şerhi"
        ]
        for c_idx, h_text in enumerate(doc_headers, start=1):
            cell = ws_docs.cell(row=2, column=c_idx, value=h_text)
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            cell.fill = cls.SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center" if c_idx in (1, 2, 4, 6, 7, 8, 9, 10) else "left", vertical="center")
            cell.border = cls.THIN_BORDER
        ws_docs.row_dimensions[2].height = 24

        r_d = 3
        for idx, doc in enumerate(documents, start=1):
            h = doc.veri.header if doc.veri else None
            kalemler = doc.veri.kalemler if doc.veri else []
            doc_qty = sum(float(k.miktar or 0.0) for k in kalemler)

            c1 = ws_docs.cell(row=r_d, column=1, value=cls.sanitize_cell_value(h.irsaliye_no if h and h.irsaliye_no else doc.id[:8]))
            c2 = ws_docs.cell(row=r_d, column=2, value=cls.sanitize_cell_value(h.duzenleme_tarihi if h and h.duzenleme_tarihi else doc.yuklenme_tarihi[:10]))
            c3 = ws_docs.cell(row=r_d, column=3, value=cls.sanitize_cell_value(h.gonderici_unvan if h and h.gonderici_unvan else "Tanımsız"))
            c4 = ws_docs.cell(row=r_d, column=4, value=cls.sanitize_cell_value(h.gonderici_vkn if h and h.gonderici_vkn else "-"))
            c5 = ws_docs.cell(row=r_d, column=5, value=cls.sanitize_cell_value(h.alici_unvan if h and h.alici_unvan else "-"))
            c6 = ws_docs.cell(row=r_d, column=6, value=cls.sanitize_cell_value(h.plaka if h and h.plaka else "-"))
            c7 = ws_docs.cell(row=r_d, column=7, value=len(kalemler))
            c8 = ws_docs.cell(row=r_d, column=8, value=round(doc_qty, 2))
            c9 = ws_docs.cell(row=r_d, column=9, value=f"{doc.veri.guven_skoru:.2f}" if doc.veri and doc.veri.guven_skoru else "-")
            c10 = ws_docs.cell(row=r_d, column=10, value=doc.durum.value)
            c11 = ws_docs.cell(row=r_d, column=11, value=cls.sanitize_cell_value(doc.veri.genel_notlar if doc.veri and doc.veri.genel_notlar else "-"))

            c1.alignment = Alignment(horizontal="center")
            c2.alignment = Alignment(horizontal="center")
            c4.alignment = Alignment(horizontal="center")
            c6.alignment = Alignment(horizontal="center")
            c7.alignment = Alignment(horizontal="center")
            c8.alignment = Alignment(horizontal="right")
            c9.alignment = Alignment(horizontal="center")
            c10.alignment = Alignment(horizontal="center")

            use_fill = cls.ZEBRA_FILL if idx % 2 == 0 else None
            for cell in (c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, c11):
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = cls.THIN_BORDER
                if use_fill:
                    cell.fill = use_fill
            ws_docs.row_dimensions[r_d].height = 20
            r_d += 1

        cls._adjust_column_widths(ws_docs)

        # Raporu kaydet
        target_file = output_path or (settings.EXPORT_DIR / f"Resmi_Mal_Kabul_Yonetici_Raporu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(target_file)
        return target_file

    @staticmethod
    def _adjust_column_widths(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len and not cell.coordinate in ws.merged_cells:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

excel_exporter = ExcelExporter()

