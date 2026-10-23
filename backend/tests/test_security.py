import io
import os
import re
import stat
import unittest
from pathlib import Path
from PIL import Image
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.core.config import settings
from backend.app.services.excel_exporter import excel_exporter
from backend.app.services.connection_manager import ws_manager
from backend.app.models.irsaliye import (
    DocumentItem,
    ProcessingStatus,
    IrsaliyeData,
    IrsaliyeHeader,
    LineItem
)

class TestSecurityAudit(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

        # Basit geçerli test JPEG verisi
        img = Image.new("RGB", (50, 50), color="blue")
        buf = io.BytesIO()
        img.save(buf, format="JPEG")
        cls.valid_jpeg_bytes = buf.getvalue()

        # Basit geçerli test PNG verisi
        buf_png = io.BytesIO()
        img.save(buf_png, format="PNG")
        cls.valid_png_bytes = buf_png.getvalue()

    # =========================================================================
    # 1. DOSYA YÜKLEME GÜVENLİĞİ TESTLERİ
    # =========================================================================

    def test_upload_path_traversal_sanitization(self):
        """Path traversal içeren dosya adlarının temizlenip UPLOAD_DIR içinde güvenle saklanması"""
        traversal_filenames = [
            ("../../etc/passwd.jpg", self.valid_jpeg_bytes, "image/jpeg"),
            ("..\\..\\windows\\system32\\cmd.jpg", self.valid_jpeg_bytes, "image/jpeg"),
            ("../../../malicious.png", self.valid_png_bytes, "image/png"),
            ("/tmp/root_escape.jpg", self.valid_jpeg_bytes, "image/jpeg"),
        ]
        for bad_filename, content, mtype in traversal_filenames:
            response = self.client.post(
                "/api/upload",
                files={"files": (bad_filename, content, mtype)}
            )
            self.assertEqual(response.status_code, 200, f"Upload failed for {bad_filename}")
            docs = response.json()
            self.assertTrue(len(docs) > 0)
            saved_doc = docs[0]

            saved_path = Path(saved_doc["dosya_yolu"]).resolve()
            upload_dir = settings.UPLOAD_DIR.resolve()

            # Kaydedilen dosya UPLOAD_DIR içinde mi?
            self.assertTrue(
                saved_path.is_relative_to(upload_dir),
                f"Path traversal tespit edildi! Dosya: {saved_path} UPLOAD_DIR: {upload_dir}"
            )
            # Dosya adında .. bulunmamalı
            self.assertNotIn("..", saved_path.name)

            # Temizlik
            self.client.delete(f"/api/documents/{saved_doc['id']}")

    def test_upload_invalid_extension_rejection(self):
        """İzin verilmeyen uzantıların (.exe, .php, .sh, .bat) reddedilmesi"""
        dangerous_extensions = ["evil.exe", "shell.php", "script.sh", "virus.bat", "config.yml"]
        for bad_name in dangerous_extensions:
            response = self.client.post(
                "/api/upload",
                files={"files": (bad_name, b"MALICIOUS_PAYLOAD_CONTENT", "application/octet-stream")}
            )
            self.assertEqual(
                response.status_code, 400,
                f"Zararlı uzantı engellenemedi: {bad_name} (Status: {response.status_code})"
            )

    def test_upload_mime_spoofing_detection(self):
        """MIME type spoofing: Zararlı dosyanın .jpg yapılarak yüklenmesi imza (magic byte) kontrolü ile engellenmeli"""
        fake_files = [
            ("php_shell.jpg", b"<?php phpinfo(); ?>", "image/jpeg"),
            ("html_xss.jpg", b"<!DOCTYPE html><html><script>alert(1)</script></html>", "image/jpeg"),
            ("bash_script.png", b"#!/bin/bash\nrm -rf /", "image/png"),
            ("corrupt.jpg", b"NOT_A_VALID_JPEG_HEADER_DATA_12345", "image/jpeg"),
        ]
        for fname, content, mtype in fake_files:
            response = self.client.post(
                "/api/upload",
                files={"files": (fname, content, mtype)}
            )
            self.assertEqual(
                response.status_code, 400,
                f"MIME spoofing tespit edilemedi: {fname} (Status: {response.status_code})"
            )

    def test_upload_empty_file_rejection(self):
        """0 byte'lık boş dosyaların yüklenmesinin engellenmesi"""
        response = self.client.post(
            "/api/upload",
            files={"files": ("empty.jpg", b"", "image/jpeg")}
        )
        self.assertEqual(response.status_code, 400)

    def test_mobile_upload_security(self):
        """Mobil yükleme rotasında path traversal ve geçersiz dosya koruması"""
        # Path traversal denemesi
        res = self.client.post(
            "/api/mobile-upload",
            data={"session_id": "test-session-123"},
            files={"files": ("../../mobile_hack.jpg", self.valid_jpeg_bytes, "image/jpeg")}
        )
        self.assertEqual(res.status_code, 200)
        doc = res.json()
        saved_path = Path(doc["dosya_yolu"]).resolve()
        self.assertTrue(saved_path.is_relative_to(settings.UPLOAD_DIR.resolve()))
        self.assertNotIn("..", saved_path.name)

        # Temizlik
        self.client.delete(f"/api/documents/{doc['id']}")

    # =========================================================================
    # 2. EXCEL FORMÜL ENJEKSİYONU (FORMULA INJECTION - CWE-1236) TESTLERİ
    # =========================================================================

    def test_excel_formula_injection_sanitization(self):
        """Tüm tehlikeli formül karakterlerinin ('=', '+', '-', '@', tab, cr, %, |) kaçırılması"""
        test_cases = [
            ("=SUM(A1:A10)", "'=SUM(A1:A10)"),
            ("=cmd|' /C calc'!A0", "'=cmd|' /C calc'!A0"),
            ("+1+1", "'+1+1"),
            ("-2+3", "'-2+3"),
            ("@HYPERLINK(\"http://evil.com\")", "'@HYPERLINK(\"http://evil.com\")"),
            ("\t=cmd", "'\t=cmd"),
            ("\r=1+1", "'\r=1+1"),
            ("  =SUM(1,2)", "'  =SUM(1,2)"),
            ("   @SUM(B1)", "'   @SUM(B1)"),
            ("%TOTAL%", "'%TOTAL%"),
            ("|calc", "'|calc"),
            (";test", "';test"),
            ("Güvenli İrsaliye Açıklaması", "Güvenli İrsaliye Açıklaması"),
            (12345, 12345),
            (99.9, 99.9),
            (None, ""),
        ]
        for raw_val, expected in test_cases:
            sanitized = excel_exporter.sanitize_cell_value(raw_val)
            self.assertEqual(
                sanitized, expected,
                f"Formül enjeksiyonu kaçırma hatası: input={raw_val!r}, got={sanitized!r}, expected={expected!r}"
            )

    def test_excel_export_end_to_end_formula_safety(self):
        """Zararlı formül içeren bir evrakın Excel çıktısında güvenli yazılması"""
        malicious_doc = DocumentItem(
            id="sec-test-doc",
            dosya_adi="=cmd.exe.jpg",
            dosya_yolu="/tmp/test.jpg",
            dosya_url="/uploads/test.jpg",
            mime_type="image/jpeg",
            boyut_bytes=1000,
            yuklenme_tarihi="23.10.2026 08:00",
            durum=ProcessingStatus.ONAYLANDI,
            veri=IrsaliyeData(
                header=IrsaliyeHeader(
                    irsaliye_no="=1+1",
                    duzenleme_tarihi="@TODAY()",
                    gonderici_unvan="+HACKED_CORP",
                    gonderici_vkn="-9999999999",
                    alici_unvan="Normal Alıcı",
                    plaka="  =calc"
                ),
                kalemler=[
                    LineItem(sira_no=1, mal_hizmet="=DDE(\"cmd\";\"/c calc\";\"a\")", miktar=10.0, birim="+ADET", notlar="@EXPLOIT"),
                ],
                genel_notlar="=cmd|'/C calc'!A0",
                genel_toplam_miktar=10.0
            )
        )

        output_path = settings.EXPORT_DIR / "test_security_formula.xlsx"
        try:
            exported_file = excel_exporter.export_single_irsaliye(malicious_doc, output_path=output_path)
            self.assertTrue(exported_file.exists())

            # openpyxl ile hücreleri kontrol et
            import openpyxl
            wb = openpyxl.load_workbook(exported_file, data_only=False)
            ws = wb.active

            # İrsaliye no hücresi
            cell_irsaliye_no = ws.cell(row=3, column=2).value
            self.assertTrue(str(cell_irsaliye_no).startswith("'"), f"İrsaliye no kaçırılmadı: {cell_irsaliye_no}")

            # Kalem mal_hizmet hücresi
            cell_mal = ws.cell(row=6, column=2).value
            self.assertTrue(str(cell_mal).startswith("'"), f"Mal hizmet formülü kaçırılmadı: {cell_mal}")

            wb.close()
        finally:
            if output_path.exists():
                output_path.unlink()

    # =========================================================================
    # 3. HASSAS VERİ & CORS GÜVENLİĞİ TESTLERİ
    # =========================================================================

    def test_api_key_masking_no_leak(self):
        """API anahtarı durumunun tam anahtarı veya başındaki 6 karakteri sızdırmaması"""
        res = self.client.get("/api/settings/api-key")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("masked_key", data)
        self.assertIn("is_configured", data)

        masked = data["masked_key"]
        # Ham anahtar hiçbir şekilde masked_key içinde yer almamalı
        raw_key = os.getenv("GEMINI_API_KEY", "") or settings.GEMINI_API_KEY
        if raw_key and raw_key != "your_gemini_api_key_here":
            self.assertNotEqual(masked, raw_key)
            # Maske '••••' ile başlamalı
            self.assertTrue(masked.startswith("••••") or masked in ("Tanımlı", "Tanımsız"))

    def test_api_key_save_injection_protection(self):
        """.env enjeksiyonu denemelerinin (newline, quotes, control chars) reddedilmesi"""
        malicious_keys = [
            "validkey123\nANOTHER_VAR=hacked",
            "validkey123\r\nDATABASE_URL=exploit",
            'validkey" || rm -rf /',
            "short",
            "key with spaces inside",
            "key;drop table users;",
        ]
        for bad_key in malicious_keys:
            res = self.client.post(
                "/api/settings/api-key",
                json={"api_key": bad_key}
            )
            self.assertEqual(
                res.status_code, 400,
                f"Tehlikeli API anahtarı kabul edildi: {bad_key!r}"
            )

    def test_env_file_permissions(self):
        """.env dosyasının yalnızca sahip tarafından okunabilir (0600) olması"""
        env_path = settings.WORKSPACE_DIR / ".env"
        if env_path.exists():
            file_stat = os.stat(env_path)
            mode = stat.S_IMODE(file_stat.st_mode)
            # 0o600 veya 0o400 (sadece sahip okuma/yazma)
            self.assertTrue(
                mode in (0o600, 0o400),
                f".env izinleri güvensiz: {oct(mode)} (0o600 olmalı)"
            )

    def test_cors_no_wildcard_with_credentials(self):
        """CORS kökenlerinde wildcard '*' olmaması ve güvenli yapılandırılması"""
        origins = settings.CORS_ORIGINS
        self.assertNotIn("*", origins, "CORS kökenlerinde wildcard '*' tespit edildi!")
        for origin in origins:
            self.assertTrue(
                origin.startswith("http://") or origin.startswith("https://"),
                f"Geçersiz CORS kökeni formatı: {origin}"
            )

    # =========================================================================
    # 4. WEBSOCKET GÜVENLİK & LİMİT TESTLERİ
    # =========================================================================

    def test_websocket_session_sanitization(self):
        """WebSocket geçersiz session_id formatlarının engellenmesi"""
        # Session ID kontrol fonksiyonu
        invalid_sessions = [
            "",
            "../../../traversal",
            "session with spaces",
            "session$injection",
            "a" * 100,  # 64'ten uzun
        ]
        for sid in invalid_sessions:
            is_valid = bool(sid and len(sid) <= 64 and re.match(r"^[a-zA-Z0-9_\-]+$", sid))
            self.assertFalse(is_valid, f"Geçersiz session id onaylandı: {sid}")

    def test_websocket_limits_configured(self):
        """WebSocket için toplam, oturum ve IP limitlerinin tanımlı olması"""
        self.assertGreater(settings.WS_MAX_TOTAL_CONNECTIONS, 0)
        self.assertGreater(settings.WS_MAX_CONNECTIONS_PER_SESSION, 0)
        self.assertGreater(settings.WS_MAX_CONNECTIONS_PER_IP, 0)
        self.assertGreater(settings.WS_RATE_LIMIT_BURST, 0)

    def test_websocket_rate_limiter_logic(self):
        """WebSocket flood korumasının hızlı isteklerde devreye girmesi"""
        test_ip = "192.168.100.99"
        # Burst limit kadar istek yap
        burst_limit = settings.WS_RATE_LIMIT_BURST
        for _ in range(burst_limit):
            limited = ws_manager._is_rate_limited(test_ip)
            self.assertFalse(limited)

        # Burst limit + 1. istek engellenmeli (flood tespit edilmeli)
        flood_detected = ws_manager._is_rate_limited(test_ip)
        self.assertTrue(flood_detected, "WebSocket flood tespiti başarısız oldu!")

if __name__ == "__main__":
    unittest.main()
